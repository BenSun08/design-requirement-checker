"""OOXML ingestion adapter: python-docx 1.2.0 + focused OOXML/XML access.

Strategy validated by the executed S1/S2 spike (docs/plan.md §4):
python-docx opens the package and provides paragraph/table/run structure,
while focused lxml access resolves what its API cannot see — the style
inheritance chain (``w:basedOn``), ``docDefaults``, merged-cell grid
geometry, tracked revisions, orphan style references and excluded
structures (content controls, text boxes, headers/footers, field codes,
footnote/endnote references, ``w:altChunk``, smart tags, nested hyperlinks).

Conservative policies enforced here (docs/constitution.md §5–§6):
- effective strike resolves run rPr -> character style chain -> paragraph
  style chain (rooted at the style marked ``w:default="1"``, not an assumed
  "Normal" id) -> docDefaults -> default off;
- unresolvable formatting (invalid values, orphan references, broken chains,
  double strike pending a product decision) stays unknown with a reason and
  is never silently converted to False;
- runs inside hyperlinks are extracted in document order via
  ``Paragraph.iter_inner_content()`` (``paragraph.runs`` omits them);
- tracked revisions, excluded structures and hidden merged-cell continuation
  content produce explicit LIMITED warnings instead of silent loss;
- merged-cell masters are extracted exactly once at their master grid
  position; vMerge continuations are skipped.

This is the only module where python-docx objects exist. It reads the file
bytes once, never writes, and returns domain values only. Offsets are Python
code-point indices with half-open ``[start, end)`` ranges.
"""

import hashlib
import io
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from docx import Document as open_docx_package
from docx.exceptions import InvalidXmlError
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.exceptions import InvalidXmlError as OxmlInvalidXmlError
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree  # type: ignore[import-untyped]

from design_requirement_checker.domain import (
    BlockType,
    Coverage,
    Document,
    DocumentBlock,
    DocumentLocation,
    Part,
    TableCellCoordinates,
    TextRun,
)

_ON = frozenset({"true", "1", "on"})
_OFF = frozenset({"false", "0", "off"})
_REVISION_TAGS = (qn("w:ins"), qn("w:del"), qn("w:moveFrom"), qn("w:moveTo"))
#: Paragraph-level structures that can hide visible text but stay outside the
#: supported extraction scope; each maps to one stable LIMITED warning token.
_STRUCTURE_WARNINGS: tuple[tuple[str, str], ...] = (
    (qn("w:instrText"), "field-code-content-excluded"),
    (qn("w:fldSimple"), "field-code-content-excluded"),
    (qn("w:footnoteReference"), "footnote-or-endnote-content-excluded"),
    (qn("w:endnoteReference"), "footnote-or-endnote-content-excluded"),
    (qn("w:smartTag"), "smart-tag-content-excluded"),
)
#: Known library-level failures of the package-open phase. lxml ships no type
#: stubs (see the type-ignore on its import); only XMLSyntaxError is used.
_OPEN_ERRORS: tuple[type[BaseException], ...] = (
    zipfile.BadZipFile,
    KeyError,
    PackageNotFoundError,
    InvalidXmlError,
    OxmlInvalidXmlError,
    etree.XMLSyntaxError,
)


class DocxReadError(Exception):
    """A .docx file could not be read into a document snapshot.

    ``reason`` is a stable token — "file-access-error" (the file could not be
    read), "invalid-or-unreadable-document" (known bad/unsupported package
    content, including the OLE container of password-protected documents) or
    "unexpected-parser-error" (an implementation bug still fails explicitly
    instead of crashing the caller); ``detail`` carries the diagnostic.
    """

    def __init__(self, reason: str, detail: str) -> None:
        super().__init__(f"{reason}: {detail}")
        self.reason = reason
        self.detail = detail


class _InvalidStrike:
    """Sentinel for ``w:strike`` values outside the ST_OnOff vocabulary."""


_INVALID_STRIKE = _InvalidStrike()


def _parse_on_off_element(rpr: Any, tag: str) -> bool | None | _InvalidStrike:
    """Read an ST_OnOff element (e.g. ``w:strike``, ``w:dstrike``) from an rPr."""
    if rpr is None:
        return None
    element = rpr.find(tag)
    if element is None:
        return None
    val = element.get(qn("w:val"))
    if val is None:
        return True  # a bare element means on
    if val in _ON:
        return True
    if val in _OFF:
        return False
    return _INVALID_STRIKE


def _strike_from_rpr(rpr: Any) -> bool | None | _InvalidStrike:
    return _parse_on_off_element(rpr, qn("w:strike"))


class _StyleResolver:
    """Style-table lookups for effective-strike chain resolution."""

    def __init__(self, doc: Any) -> None:
        styles_element = doc.styles.element
        self.style_map: dict[str, Any] = {
            style.get(qn("w:styleId")): style for style in styles_element.findall(qn("w:style"))
        }
        doc_defaults = styles_element.find(qn("w:docDefaults"))
        rpr_default = doc_defaults.find(qn("w:rPrDefault")) if doc_defaults is not None else None
        self.doc_defaults_strike: bool | None | _InvalidStrike = _strike_from_rpr(
            rpr_default.find(qn("w:rPr")) if rpr_default is not None else None
        )
        # The default paragraph style is the one marked w:default="1" — its
        # styleId is not necessarily "Normal" in enterprise templates. When the
        # marker is absent or ambiguous, no default style is fabricated.
        default_styles = [
            style
            for style in styles_element.findall(qn("w:style"))
            if style.get(qn("w:type")) == "paragraph"
            and style.get(qn("w:default")) == "1"
            and style.get(qn("w:styleId")) is not None
        ]
        self.default_paragraph_style_id: str | None = (
            default_styles[0].get(qn("w:styleId")) if len(default_styles) == 1 else None
        )

    def chain_strike(self, style_id: str) -> tuple[bool | None, str]:
        """Walk a basedOn chain; returns (value, "") or (None, unknown_reason)."""
        visited: set[str] = set()
        current: Any = self.style_map.get(style_id)
        if current is None:
            return (None, "orphan-style-reference")
        while current is not None:
            sid: str | None = current.get(qn("w:styleId"))
            if sid is None or sid in visited:
                return (None, "style-chain-cycle")
            visited.add(sid)
            value = _strike_from_rpr(current.find(qn("w:rPr")))
            if isinstance(value, _InvalidStrike):
                return (None, "invalid-strike-value")
            if value is not None:
                return (value, "")
            based_on = current.find(qn("w:basedOn"))
            if based_on is None:
                return (None, "")
            current = self.style_map.get(based_on.get(qn("w:val")))
            if current is None:
                return (None, "broken-style-chain")
        return (None, "")


def _read_direct_strike(run: Run) -> tuple[bool | None, str]:
    """Read the run-level strike from its rPr element; no library exceptions.

    ``run.font.strike`` raises ``InvalidXmlError`` for values outside the
    ST_OnOff vocabulary; the element-level read classifies the same cases as
    an unknown value with a reason instead of catching broad exceptions.
    """
    value = _strike_from_rpr(run._r.find(qn("w:rPr")))
    if isinstance(value, _InvalidStrike):
        return (None, "invalid-strike-value")
    return (value, "")


def _read_double_strike(run: Run) -> bool | None:
    """Element-level ``w:dstrike`` read; invalid values count as absent."""
    value = _parse_on_off_element(run._r.find(qn("w:rPr")), qn("w:dstrike"))
    if isinstance(value, _InvalidStrike):
        return None
    return value


def _resolve_strike(
    run: Run, paragraph_style_id: str | None, resolver: _StyleResolver, double: bool | None
) -> tuple[bool | None, str, str]:
    """Return (effective_strike, origin, unknown_reason)."""
    direct, direct_reason = _read_direct_strike(run)
    if direct_reason:
        return (None, "direct", direct_reason)
    value: bool | None = None
    origin = "default-off"
    if direct is not None:
        value, origin = direct, "direct"
    else:
        rpr = run._r.find(qn("w:rPr"))
        rstyle = rpr.find(qn("w:rStyle")) if rpr is not None else None
        if rstyle is not None:
            char_value, char_reason = resolver.chain_strike(rstyle.get(qn("w:val")))
            if char_reason:
                return (None, "character-style", char_reason)
            if char_value is not None:
                value, origin = char_value, "character-style"
        if value is None and paragraph_style_id is not None:
            para_value, para_reason = resolver.chain_strike(paragraph_style_id)
            if para_reason:
                return (None, "paragraph-style", para_reason)
            if para_value is not None:
                value, origin = para_value, "paragraph-style"
        if value is None:
            doc_default = resolver.doc_defaults_strike
            if isinstance(doc_default, _InvalidStrike):
                return (None, "doc-defaults", "invalid-strike-value")
            if doc_default is not None:
                value, origin = doc_default, "doc-defaults"
    if value is None:
        # The chain resolved without defining strike: OOXML defaults it to off.
        value = False
    # Double strike is reliably detectable, but whether it counts as deletion
    # formatting is an open product question. Conservative: unknown unless
    # single strike is explicitly True.
    if double is True and value is not True:
        return (None, origin, "double-strike")
    return (value, origin, "")


def _paragraph_style_id(paragraph: Paragraph, resolver: _StyleResolver) -> str | None:
    ppr = paragraph._p.find(qn("w:pPr"))
    pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
    if pstyle is not None:
        style_id: str | None = pstyle.get(qn("w:val"))
        return style_id
    return resolver.default_paragraph_style_id


def _iter_paragraph_runs(paragraph: Paragraph) -> Iterator[Run]:
    """Yield visible runs in document order, including hyperlink-wrapped runs.

    ``paragraph.runs`` omits runs inside ``w:hyperlink``; python-docx 1.2.0
    exposes the paragraph's inner content (Run | Hyperlink) in document order,
    and ``Hyperlink.runs`` carries the link's own runs with their formatting.
    The link URL itself is never document text and is not yielded.
    """
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            yield from item.runs
        else:
            yield item


def _block_id(
    part: Part, cell: TableCellCoordinates | None, ancestor_path: tuple[TableCellCoordinates, ...]
) -> str:
    cells = [*ancestor_path]
    if cell is not None:
        cells.append(cell)
    prefix = ">".join(f"t{c.table_index}r{c.row_index}c{c.column_index}" for c in cells)
    return f"{prefix + ':' if prefix else ''}{part}:p"


def _detect_paragraph_exclusions(p_el: Any, warnings: list[str]) -> None:
    """Warn about paragraph content that stays outside the supported scope."""
    if any(p_el.findall(f".//{tag}") for tag in _REVISION_TAGS):
        warnings.append("tracked-revisions-unsupported")
    if p_el.findall(f".//{qn('w:txbxContent')}"):
        warnings.append("textbox-content-excluded")
    if p_el.findall(f".//{qn('w:sdt')}"):
        warnings.append("content-control-content-excluded")
    for tag, token in _STRUCTURE_WARNINGS:
        if p_el.findall(f".//{tag}"):
            warnings.append(token)
    if p_el.findall(f".//{qn('w:hyperlink')}//{qn('w:hyperlink')}"):
        # Invalid OOXML: a hyperlink nested inside another hyperlink. Its runs
        # are invisible to every python-docx run accessor; exclude explicitly.
        warnings.append("nested-hyperlink-content-excluded")


def _collect_paragraph(
    document_id: str,
    paragraph: Paragraph,
    part: Part,
    cell: TableCellCoordinates | None,
    ancestor_path: tuple[TableCellCoordinates, ...],
    paragraph_index: int,
    resolver: _StyleResolver,
    blocks: list[DocumentBlock],
    warnings: list[str],
) -> None:
    p_el: Any = paragraph._p
    _detect_paragraph_exclusions(p_el, warnings)
    style_id = _paragraph_style_id(paragraph, resolver)
    runs: list[TextRun] = []
    offset = 0
    for run in _iter_paragraph_runs(paragraph):
        text: str = run.text
        if not text:
            continue  # e.g. w:fldChar / w:instrText carrier runs carry no text
        double = _read_double_strike(run)
        strike, origin, reason = _resolve_strike(run, style_id, resolver, double)
        runs.append(
            TextRun(
                text=text,
                start_offset=offset,
                end_offset=offset + len(text),
                effective_strike=strike,
                strike_origin=origin,
                strike_reason=reason,
                double_strike=double,
            )
        )
        offset += len(text)
    block_type = BlockType.PARAGRAPH if part == "body" else BlockType.TABLE_CELL_PARAGRAPH
    block_id = _block_id(part, cell, ancestor_path) + str(paragraph_index)
    location = DocumentLocation(
        document_id=document_id,
        block_id=block_id,
        block_type=block_type,
        part=part,
        paragraph_index=paragraph_index,
        cell=cell,
        ancestor_path=ancestor_path,
    )
    blocks.append(
        DocumentBlock(
            block_id=block_id,
            block_type=block_type,
            text="".join(run.text for run in runs),
            runs=tuple(runs),
            location=location,
        )
    )


def _has_visible_text(element: Any) -> bool:
    """True when any ``w:t`` under the element carries non-whitespace text."""
    return any((node.text or "").strip() != "" for node in element.findall(f".//{qn('w:t')}"))


def _collect_table(
    table: Table,
    ancestor_path: tuple[TableCellCoordinates, ...],
    table_index: int,
    resolver: _StyleResolver,
    document_id: str,
    blocks: list[DocumentBlock],
    warnings: list[str],
) -> None:
    """Traverse ``w:tc`` elements directly (``row.cells`` duplicates merged cells).

    Merged-cell convention: a gridSpan master is extracted once at its starting
    grid column; a vMerge master is extracted once at its (row, col); vMerge
    continuation cells are skipped, and non-empty continuation content raises
    an explicit warning because Word does not display it.

    Content controls (``w:sdt``) can also wrap an entire table row (Word
    repeating-section controls) or a single cell. They stay outside the
    supported scope like body/cell-level content controls and must raise the
    same explicit warning — their text never vanishes silently.
    """
    row_index = 0
    for tbl_child in table._tbl:
        if tbl_child.tag == qn("w:sdt"):
            warnings.append("content-control-content-excluded")
            continue
        if tbl_child.tag != qn("w:tr"):
            continue
        column_cursor = 0
        for row_child in tbl_child:
            if row_child.tag == qn("w:sdt"):
                warnings.append("content-control-content-excluded")
                continue
            if row_child.tag != qn("w:tc"):
                continue
            tc = row_child
            tc_pr = tc.find(qn("w:tcPr"))
            grid_span = 1
            v_merge = None
            if tc_pr is not None:
                span_el = tc_pr.find(qn("w:gridSpan"))
                if span_el is not None:
                    span_value: str | None = span_el.get(qn("w:val"))
                    grid_span = int(span_value) if span_value else 1
                v_merge = tc_pr.find(qn("w:vMerge"))
            is_continuation = v_merge is not None and (
                v_merge.get(qn("w:val")) in (None, "continue")
            )
            if is_continuation:
                if _has_visible_text(tc):
                    warnings.append("merged-cell-continuation-content-excluded")
                column_cursor += grid_span
                continue
            cell = TableCellCoordinates(
                table_index=table_index, row_index=row_index, column_index=column_cursor
            )
            cell_parent = _Cell(tc, table)
            paragraph_index = 0
            nested_index = 0
            for child in tc:
                if child.tag == qn("w:p"):
                    _collect_paragraph(
                        document_id,
                        Paragraph(child, cell_parent),
                        "table-cell",
                        cell,
                        ancestor_path,
                        paragraph_index,
                        resolver,
                        blocks,
                        warnings,
                    )
                    paragraph_index += 1
                elif child.tag == qn("w:tbl"):
                    _collect_table(
                        Table(child, cell_parent),
                        (*ancestor_path, cell),
                        nested_index,
                        resolver,
                        document_id,
                        blocks,
                        warnings,
                    )
                    nested_index += 1
                elif child.tag == qn("w:sdt"):
                    warnings.append("content-control-content-excluded")
            column_cursor += grid_span
        row_index += 1


def _collect_body(
    doc: Any,
    document_id: str,
    resolver: _StyleResolver,
    blocks: list[DocumentBlock],
    warnings: list[str],
) -> None:
    body = doc.element.body
    if body.findall(f".//{qn('w:altChunk')}"):
        # w:altChunk imports external content at render time; it is never
        # extracted and must not allow a silent COMPLETE result.
        warnings.append("alt-chunk-content-excluded")
    paragraph_index = 0
    table_index = 0
    for child in body:
        if child.tag == qn("w:p"):
            _collect_paragraph(
                document_id,
                Paragraph(child, doc._body),
                "body",
                None,
                (),
                paragraph_index,
                resolver,
                blocks,
                warnings,
            )
            paragraph_index += 1
        elif child.tag == qn("w:tbl"):
            _collect_table(
                Table(child, doc._body), (), table_index, resolver, document_id, blocks, warnings
            )
            table_index += 1
        elif child.tag == qn("w:sdt"):
            warnings.append("content-control-content-excluded")


def _collect_header_footer_warnings(doc: Any, warnings: list[str]) -> None:
    """Detect content in every header/footer variant of every section.

    All variants (default, first-page, even-page) share one stable token per
    side ("header-content-not-checked" / "footer-content-not-checked"); both
    paragraph and table content count. Variant parts stay outside the checked
    body scope. The linked-to-previous guard must run first: reading content
    of a linked container would create a part in the in-memory package.
    """
    for section in doc.sections:
        variants = (
            (section.header, "header"),
            (section.first_page_header, "header"),
            (section.even_page_header, "header"),
            (section.footer, "footer"),
            (section.first_page_footer, "footer"),
            (section.even_page_footer, "footer"),
        )
        for container, label in variants:
            if container.is_linked_to_previous:
                continue  # no explicit part of its own in this section
            has_text = any(paragraph.text.strip() for paragraph in container.paragraphs)
            has_table_text = any(_has_visible_text(table._tbl) for table in container.tables)
            if has_text or has_table_text:
                warnings.append(f"{label}-content-not-checked")


def read_document(path: Path) -> Document:
    """Read a .docx file without modifying it; raise DocxReadError on failure."""
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise DocxReadError("file-access-error", f"{type(exc).__name__}: {exc}") from exc
    fingerprint = hashlib.sha256(data).hexdigest()
    document_id = f"sha256:{fingerprint}"
    try:
        package = open_docx_package(io.BytesIO(data))
    except _OPEN_ERRORS as exc:
        # BadZipFile also covers the OLE compound-file container Word produces
        # for password-protected documents (it is not a ZIP package).
        raise DocxReadError(
            "invalid-or-unreadable-document", f"{type(exc).__name__}: {exc}"[:200]
        ) from exc
    except Exception as exc:
        # Unexpected library failure at the open boundary: fail explicitly
        # instead of crashing the caller, under its own category.
        raise DocxReadError(
            "unexpected-parser-error", f"{type(exc).__name__}: {exc}"[:200]
        ) from exc
    try:
        warnings: list[str] = []
        blocks: list[DocumentBlock] = []
        resolver = _StyleResolver(package)
        _collect_body(package, document_id, resolver, blocks, warnings)
        _collect_header_footer_warnings(package, warnings)
    except Exception as exc:
        # Boundary guard: an implementation bug in collection must not crash
        # the UI, and it is not classified as an invalid document.
        raise DocxReadError(
            "unexpected-parser-error", f"{type(exc).__name__}: {exc}"[:200]
        ) from exc
    unique_warnings = sorted(set(warnings))
    coverage = Coverage.LIMITED if unique_warnings else Coverage.COMPLETE
    return Document(
        document_id=document_id,
        filename=path.name,
        content_fingerprint=fingerprint,
        blocks=tuple(blocks),
        coverage=coverage,
        warnings=tuple(unique_warnings),
    )
