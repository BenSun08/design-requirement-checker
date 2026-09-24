"""Deterministic synthetic DOCX fixtures for the S1/S2 technical spike.

Each builder writes a fixed-content fixture to ``path``. Expected outcomes for
these fixtures are hand-written as independent labels inside the spike test
files; do not derive expectations from probe output.

Cases that python-docx cannot generate through its public API are created by
small, documented raw-OOXML patches (double strike, invalid strike value,
orphan style reference, docDefaults strike, tracked revisions, content
controls, text boxes, hidden vertically-merged-cell content).
"""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

NSDECL = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _raw(xml: str):  # type: ignore[no-untyped-def]
    return parse_xml(xml)


def build_normal(path: Path) -> Path:
    """Body paragraphs, split runs, adjacent same-format runs, whitespace, empty paragraph."""
    doc = Document()
    doc.add_paragraph("plain body paragraph one")
    p1 = doc.add_paragraph()
    p1.add_run("Hello ")
    bold = p1.add_run("bold")
    bold.bold = True
    p1.add_run(" world")
    p2 = doc.add_paragraph()
    p2.add_run("Alpha")
    p2.add_run("Beta")
    doc.add_paragraph("A  B\tC")
    doc.add_paragraph("")
    doc.save(path)
    return path


def build_table(path: Path) -> Path:
    """One body paragraph, a 2x2 table (one cell with two paragraphs), one trailing paragraph."""
    doc = Document()
    doc.add_paragraph("intro")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "r0c0"
    table.cell(0, 1).text = "r0c1"
    table.cell(1, 0).text = "r1c0"
    cell = table.cell(1, 1)
    cell.text = "cell para A"
    cell.add_paragraph("cell para B")
    doc.add_paragraph("outro")
    doc.save(path)
    return path


def build_table_content_controls(path: Path) -> Path:
    """A 2x2 table whose row-2 second cell is wrapped in a ``w:sdt`` content
    control, plus a whole extra row wrapped in a ``w:sdt`` at the table level
    (Word repeating-section layout). Word renders both; the declared supported
    scope excludes them, so they must raise an explicit warning instead of
    vanishing silently. The unwrapped cells must still be extracted."""
    doc = Document()
    doc.add_paragraph("intro")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "VISIBLE-A"
    table.cell(0, 1).text = "VISIBLE-B"
    table.cell(1, 0).text = "VISIBLE-C"
    tr = table.rows[1]._tr
    tc = tr.tc_lst[1]
    sdt_cell = _raw(
        f'<w:sdt {NSDECL}><w:sdtPr><w:id w:val="7"/></w:sdtPr><w:sdtContent>'
        "<w:tc><w:p><w:r><w:t>SDT-CELL</w:t></w:r></w:p></w:tc></w:sdtContent></w:sdt>"
    )
    tr.replace(tc, sdt_cell)
    table._tbl.append(
        _raw(
            f'<w:sdt {NSDECL}><w:sdtPr><w:id w:val="8"/></w:sdtPr><w:sdtContent>'
            "<w:tr><w:tc><w:p><w:r><w:t>SDT-ROW-A</w:t></w:r></w:p></w:tc>"
            "<w:tc><w:p><w:r><w:t>SDT-ROW-B</w:t></w:r></w:p></w:tc></w:tr>"
            "</w:sdtContent></w:sdt>"
        )
    )
    doc.save(path)
    return path


def build_table_long_cell(path: Path) -> Path:
    """One body paragraph and a 1x2 table whose second cell carries a long
    multi-paragraph 功能描述 — the real-document pattern where requirement
    text lives inside table cells and must stay visible in the document
    inspection view."""
    doc = Document()
    doc.add_paragraph("功能要求汇总")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "延时上电"
    long_line = (
        "KL30电后延时3s输出，检测到电压小于9V时立即关闭输出；"
        "具有过压保护功能，过压阈值为16V，过压后延时500ms恢复。"
    )
    cell = table.cell(0, 1)
    cell.text = long_line
    cell.add_paragraph("输出具有短路保护功能，短路解除后自动恢复。")
    doc.save(path)
    return path


def build_nested_merged(path: Path) -> Path:
    """3x3 grid with a gridSpan master, a vMerge master, patched hidden continuation
    content, and a nested table inside a cell."""
    doc = Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "H-MERGE"
    table.cell(0, 2).text = "r0c2"
    table.cell(1, 0).text = "r1c0"
    table.cell(1, 1).text = "r1c1"
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(1, 2).text = "V-MERGE"
    table.cell(2, 0).text = "r2c0"
    outer = table.cell(2, 1)
    outer.text = "outer-para"
    nested = outer.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "NESTED-CELL"
    # Raw patch: the vertically merged continuation cell at (row 2, grid col 2)
    # contains text that Word does not display. Word requires each tc to end
    # with a paragraph; the continuation tc keeps an empty w:p which we fill.
    continuation_tc = table.rows[2]._tr.tc_lst[2]
    continuation_p = continuation_tc.find(qn("w:p"))
    continuation_p.append(_raw(f"<w:r {NSDECL}><w:t>HIDDEN-CONT</w:t></w:r>"))
    doc.save(path)
    return path


def build_strike_matrix(path: Path) -> Path:
    """S1 formatting matrix in one document."""
    doc = Document()
    doc.add_paragraph().add_run("direct strike").font.strike = True
    doc.add_paragraph().add_run("explicit false").font.strike = False
    p2 = doc.add_paragraph()
    p2.add_run("keep ")
    p2.add_run("gone ").font.strike = True
    p2.add_run("tail")
    struck = doc.styles.add_style("StruckStyle", WD_STYLE_TYPE.PARAGRAPH)
    struck.base_style = doc.styles["Normal"]
    struck.font.strike = True
    doc.add_paragraph("inherited", style=struck)
    p4 = doc.add_paragraph(style=struck)
    p4.add_run("override").font.strike = False
    # Raw patch: double strikethrough (python-docx cannot set dstrike via font).
    p5 = doc.add_paragraph()
    r5 = p5.add_run("dbl")
    r5._r.get_or_add_rPr().append(_raw(f"<w:dstrike {NSDECL}/>"))
    # Raw patch: invalid strike value, outside the ST_OnOff vocabulary.
    p6 = doc.add_paragraph()
    r6 = p6.add_run("maybe")
    r6._r.get_or_add_rPr().append(_raw(f'<w:strike {NSDECL} w:val="maybe"/>'))
    # Raw patch: character style reference to a style that does not exist.
    p7 = doc.add_paragraph()
    r7 = p7.add_run("orphan")
    r7._r.get_or_add_rPr().append(_raw(f'<w:rStyle {NSDECL} w:val="MissingStyle"/>'))
    doc.save(path)
    return path


def build_docdefaults_strike(path: Path) -> Path:
    """docDefaults-level strike; one plain paragraph, one explicit-false override."""
    doc = Document()
    doc_defaults_rpr = (
        doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    )  # type: ignore[union-attr]
    doc_defaults_rpr.append(_raw(f"<w:strike {NSDECL}/>"))
    doc.add_paragraph("plain")
    p1 = doc.add_paragraph()
    p1.add_run("off").font.strike = False
    doc.save(path)
    return path


def build_tracked_revisions(path: Path) -> Path:
    """One paragraph with tracked insert/delete plus one clean paragraph."""
    doc = Document()
    p0 = doc.add_paragraph()
    p0.add_run("kept ")
    # Raw patch: tracked insertion and deletion (python-docx has no revision API).
    p0._p.append(
        _raw(
            f'<w:ins {NSDECL} w:id="1" w:author="spike" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:t>INSERTED</w:t></w:r></w:ins>"
        )
    )
    p0._p.append(
        _raw(
            f'<w:del {NSDECL} w:id="2" w:author="spike" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:delText>REMOVED</w:delText></w:r></w:del>"
        )
    )
    doc.add_paragraph("clean paragraph")
    doc.save(path)
    return path


def build_excluded_parts(path: Path) -> Path:
    """Header/footer text, a body-level content-control paragraph, and a text box."""
    doc = Document()
    doc.add_paragraph("visible body")
    section = doc.sections[0]
    section.header.paragraphs[0].text = "HEADER TEXT"
    section.footer.paragraphs[0].text = "FOOTER TEXT"
    # Raw patch: body-level content control (w:sdt) holding a paragraph.
    # Inserted before sectPr to keep the required body element order.
    sect_pr = doc.element.body.find(qn("w:sectPr"))
    sect_pr.addprevious(
        _raw(
            f'<w:sdt {NSDECL}><w:sdtPr><w:id w:val="1"/></w:sdtPr>'
            "<w:sdtContent><w:p><w:r><w:t>SDT PARA</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        )
    )
    # Raw patch: minimal text box (drawing/txbxContent) inside a body paragraph run.
    p1 = doc.add_paragraph()
    p1.add_run("host ")
    textbox_run = p1.add_run()
    textbox_run._r.append(
        _raw(
            "<mc:AlternateContent "
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            '<mc:Choice Requires="wps">'
            "<w:drawing "
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            f'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" {NSDECL}>'
            '<wp:inline><a:graphic><a:graphicData uri="wordprocessingShape">'
            "<wps:wsp><wps:txbx><w:txbxContent>"
            "<w:p><w:r><w:t>TEXTBOX TEXT</w:t></w:r></w:p>"
            "</w:txbxContent></wps:txbx></wps:wsp>"
            "</a:graphicData></a:graphic></wp:inline></w:drawing>"
            "</mc:Choice></mc:AlternateContent>"
        )
    )
    doc.save(path)
    return path


def build_hyperlink(path: Path) -> Path:
    """Hyperlink-wrapped runs (w:anchor links need no relationships):
    prefix/link/suffix order with a struck link run, a multi-run link, a link
    inside a table cell, and an invalid nested hyperlink whose inner run is
    invisible to paragraph.runs, iter_inner_content and Hyperlink.runs."""
    doc = Document()
    p0 = doc.add_paragraph()
    p0.add_run("Prefix ")
    p0._p.append(
        _raw(
            f'<w:hyperlink {NSDECL} w:anchor="bm1">'
            "<w:r><w:rPr><w:strike/></w:rPr><w:t>hyperlink text</w:t></w:r>"
            "</w:hyperlink>"
        )
    )
    p0.add_run(" suffix")
    p1 = doc.add_paragraph()
    p1._p.append(
        _raw(
            f'<w:hyperlink {NSDECL} w:anchor="bm2">'
            "<w:r><w:t>linkA</w:t></w:r>"
            "<w:r><w:rPr><w:strike/></w:rPr><w:t>linkB</w:t></w:r>"
            "</w:hyperlink>"
        )
    )
    cell_paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    cell_paragraph._p.append(
        _raw(f'<w:hyperlink {NSDECL} w:anchor="bm3"><w:r><w:t>cell link</w:t></w:r></w:hyperlink>')
    )
    p2 = doc.add_paragraph()
    p2._p.append(
        _raw(
            f'<w:hyperlink {NSDECL} w:anchor="bm4">'
            "<w:r><w:t>outer</w:t></w:r>"
            f'<w:hyperlink {NSDECL} w:anchor="bm5"><w:r><w:t>nested</w:t></w:r></w:hyperlink>'
            "</w:hyperlink>"
        )
    )
    doc.save(path)
    return path


def build_header_variants(path: Path) -> Path:
    """Default and first-page headers with paragraph text, plus an even-page
    footer whose only meaningful content is a table cell."""
    doc = Document()
    doc.add_paragraph("visible body")
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    doc.settings.odd_and_even_pages_header_footer = True
    section.header.paragraphs[0].add_run("DEFAULT HEADER")
    section.first_page_header.paragraphs[0].add_run("FIRST PAGE HEADER")
    even_footer = section.even_page_footer
    footer_table = even_footer.add_table(1, 1, Inches(2))
    footer_table.cell(0, 0).paragraphs[0].add_run("EVEN FOOTER TABLE CELL")
    doc.save(path)
    return path


def build_custom_default_style(path: Path) -> Path:
    """The default paragraph style is 'CorpBody' (w:default='1'), not 'Normal';
    'Normal' keeps no strike while CorpBody carries one, so paragraphs without
    an explicit pStyle must resolve strike through CorpBody."""
    doc = Document()
    styles = doc.styles.element
    normal = next(
        style for style in styles.findall(qn("w:style")) if style.get(qn("w:styleId")) == "Normal"
    )
    normal.attrib.pop(qn("w:default"), None)
    styles.append(
        _raw(
            f'<w:style {NSDECL} w:type="paragraph" w:default="1" w:styleId="CorpBody">'
            '<w:name w:val="CorpBody"/><w:qFormat/>'
            "<w:rPr><w:strike/></w:rPr>"
            "</w:style>"
        )
    )
    doc.add_paragraph("plain inherits CorpBody strike")
    p1 = doc.add_paragraph()
    p1.add_run("explicit off").font.strike = False
    doc.save(path)
    return path


def build_unsupported_structures(path: Path) -> Path:
    """Field codes (complex w:instrText and simple w:fldSimple), footnote and
    endnote references, a body-level w:altChunk and a w:smartTag wrapping a
    run — content that stays outside the supported scope and must never allow
    a silent COMPLETE result."""
    doc = Document()
    doc.add_paragraph("intro")
    p_field = doc.add_paragraph()
    for fragment in (
        f'<w:r {NSDECL}><w:fldChar w:fldCharType="begin"/></w:r>',
        f'<w:r {NSDECL}><w:instrText xml:space="preserve"> REF _Ref1 \\h </w:instrText></w:r>',
        f'<w:r {NSDECL}><w:fldChar w:fldCharType="separate"/></w:r>',
        f"<w:r {NSDECL}><w:t>field result text</w:t></w:r>",
        f'<w:r {NSDECL}><w:fldChar w:fldCharType="end"/></w:r>',
    ):
        p_field._p.append(_raw(fragment))
    p_simple = doc.add_paragraph()
    p_simple._p.append(
        _raw(
            f'<w:fldSimple {NSDECL} w:instr=" REF _Ref2 \\h ">'
            "<w:r><w:t>fldSimple cached result</w:t></w:r>"
            "</w:fldSimple>"
        )
    )
    p_note = doc.add_paragraph()
    p_note.add_run("body with footnote")
    p_note._p.append(_raw(f'<w:r {NSDECL}><w:footnoteReference w:id="2"/></w:r>'))
    p_end = doc.add_paragraph()
    p_end.add_run("body with endnote")
    p_end._p.append(_raw(f'<w:r {NSDECL}><w:endnoteReference w:id="2"/></w:r>'))
    p_tag = doc.add_paragraph()
    p_tag._p.append(
        _raw(
            f'<w:smartTag {NSDECL} w:element="city">'
            "<w:r><w:t>smart tag text</w:t></w:r>"
            "</w:smartTag>"
        )
    )
    p_tag.add_run(" tail")
    # altChunk at body level, before sectPr, referencing an absent relationship
    # (the package opens fine; only the reference is kept for detection).
    sect_pr = doc.element.body.find(qn("w:sectPr"))
    sect_pr.addprevious(
        _raw(
            f'<w:altChunk {NSDECL} r:id="rId100" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
        )
    )
    doc.add_paragraph("clean paragraph")
    doc.save(path)
    return path


def build_ole_container(path: Path) -> Path:
    """The 512-byte OLE compound-file header Word produces for password-
    protected documents (signature D0 CF 11 E0 A1 B1 1A E1). Not a ZIP/OOXML
    package; exercises the container-format failure path only — not a full
    encrypted-document validation."""
    header = bytearray(512)
    header[0:8] = bytes.fromhex("d0cf11e0a1b11ae1")
    header[24:26] = (0x003E).to_bytes(2, "little")  # minor version
    header[26:28] = (0x0003).to_bytes(2, "little")  # major version
    header[28:30] = (0xFFFE).to_bytes(2, "little")  # little-endian byte order
    header[30:32] = (0x0009).to_bytes(2, "little")  # 512-byte sectors
    header[32:34] = (0x0006).to_bytes(2, "little")  # 64-byte mini sectors
    path.write_bytes(bytes(header))
    return path


def build_requirement_strike(path: Path) -> Path:
    """S6 composition fixture: three requirement sentences, one per strike row.

    p0 fully active (CONFIGURED row), p1 with only the value part struck
    (partial-strike row), p2 fully struck (STRUCK_OUT row). Detection phrases
    differ per paragraph so each S6 item matches exactly one paragraph.
    """
    doc = Document()
    doc.add_paragraph("A功能增加延时3s")
    p1 = doc.add_paragraph()
    p1.add_run("B功能增加延时")
    struck = p1.add_run("3s")
    struck.font.strike = True
    p2 = doc.add_paragraph("C功能增加延时2s")
    for run in p2.runs:
        run.font.strike = True
    doc.save(path)
    return path


def build_empty(path: Path) -> Path:
    """A valid document with no paragraphs and no tables."""
    doc = Document()
    doc.save(path)
    return path


def build_malformed(path: Path) -> Path:
    """Bytes that are not an OOXML package at all."""
    path.write_bytes(b"this is not a zip package")
    return path
