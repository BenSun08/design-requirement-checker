"""Task 6 T6.2 — evaluator harness tests.

The corpus here is SYNTHETIC and exists only to prove the harness behaves
correctly. Synthetic fixtures are NOT historical validation evidence and must
never be presented as measured Task 6 results (see validation/README.md,
ground-truth separation rule).
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document

from design_requirement_checker.baseline_store import save_baseline
from design_requirement_checker.domain import CheckItem
from validation.harness import run_validation

# ---------------------------------------------------------------------------
# Synthetic corpus construction helpers
# ---------------------------------------------------------------------------


def _build_docx(path: Path) -> None:
    """One active 3s paragraph (DIFFERENT vs expected 2s), one fully struck
    paragraph, and no mention of the third function."""
    doc = Document()
    doc.add_paragraph("A功能增加延时3s")
    struck = doc.add_paragraph("B功能增加延时2s")
    for run in struck.runs:
        run.font.strike = True
    doc.save(path)


def _items() -> tuple[CheckItem, ...]:
    return (
        CheckItem(
            item_id="item-a",
            code="DR-A",
            name="A功能",
            detection_phrase="A功能增加延时",
            expected_description="A功能增加延时2s",
        ),
        CheckItem(
            item_id="item-b",
            code="DR-B",
            name="B功能",
            detection_phrase="B功能增加延时",
        ),
        CheckItem(
            item_id="item-c",
            code="DR-C",
            name="C功能",
            detection_phrase="C功能增加延时",
        ),
    )


def _labels_doc() -> dict[str, object]:
    return {
        "schemaVersion": 1,
        "documentId": "hist-001",
        "documentFile": "corpus/hist-001.docx",
        "labelledBy": "synthetic-test-fixture",
        "sanitizationState": "synthetic",
        "notes": "synthetic evaluator self-test corpus, not historical evidence",
        "labels": [
            {
                "itemId": "item-a",
                "expectedResolution": "RESOLVED",
                "expectedStatus": "CONFIGURED",
                "expectedComparison": "DIFFERENT",
                "expectedStrike": "NONE",
                "includeInMetrics": True,
            },
            {
                "itemId": "item-b",
                "expectedResolution": "RESOLVED",
                "expectedStatus": "STRUCK_OUT",
                "expectedStrike": "FULL",
                "includeInMetrics": True,
            },
            {
                "itemId": "item-c",
                "expectedResolution": "RESOLVED",
                "expectedStatus": "MISSING",
                "includeInMetrics": True,
            },
        ],
    }


@pytest.fixture()
def corpus(tmp_path: Path) -> Path:
    """Lay out a valid synthetic corpus; return the manifest path."""
    corpus_dir = tmp_path / "corpus"
    corpus_dir.mkdir()
    labels_dir = tmp_path / "labels"
    labels_dir.mkdir()
    _build_docx(corpus_dir / "hist-001.docx")
    save_baseline(corpus_dir / "baseline.json", _items(), "baseline-synthetic")
    (labels_dir / "hist-001.json").write_text(
        json.dumps(_labels_doc(), ensure_ascii=False), encoding="utf-8"
    )
    manifest = {
        "schemaVersion": 1,
        "baselineFile": "corpus/baseline.json",
        "documents": [
            {
                "documentId": "hist-001",
                "documentFile": "corpus/hist-001.docx",
                "labelsFile": "labels/hist-001.json",
            }
        ],
    }
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    return manifest_path


# ---------------------------------------------------------------------------
# Harness
# ---------------------------------------------------------------------------


def test_harness_produces_expected_case_records(corpus: Path) -> None:
    results = run_validation(corpus)
    cases = {c["itemId"]: c for c in results["cases"]}
    assert set(cases) == {"item-a", "item-b", "item-c"}

    case_a = cases["item-a"]
    assert case_a["observed"]["resolution"] == "RESOLVED"
    assert case_a["observed"]["status"] == "CONFIGURED"
    assert case_a["observed"]["comparison"] == "DIFFERENT"
    assert case_a["correct"] == {
        "resolution": True,
        "status": True,
        "comparison": True,
        "strike": True,
    }

    case_b = cases["item-b"]
    assert case_b["observed"]["status"] == "STRUCK_OUT"
    assert case_b["observed"]["strike"] == "FULL"
    assert case_b["correct"]["status"] is True

    case_c = cases["item-c"]
    assert case_c["observed"]["status"] == "MISSING"
    assert case_c["correct"]["status"] is True


def test_harness_records_discrepancy(corpus: Path) -> None:
    """A wrong label must be detected as incorrect, not silently accepted."""
    labels_path = corpus.parent / "labels" / "hist-001.json"
    data = json.loads(labels_path.read_text(encoding="utf-8"))
    for label in data["labels"]:
        if label["itemId"] == "item-c":
            label["expectedStatus"] = "CONFIGURED"  # deliberately wrong
    labels_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    results = run_validation(corpus)
    case_c = next(c for c in results["cases"] if c["itemId"] == "item-c")
    assert case_c["observed"]["status"] == "MISSING"
    assert case_c["correct"]["status"] is False
    assert case_c["discrepancyCategory"] == ""


def test_harness_reproducibility_metadata(corpus: Path) -> None:
    metadata = run_validation(corpus)["metadata"]
    assert metadata["ruleRevision"] == "task3-v1"
    assert metadata["validationSchemaVersion"] == 1
    assert metadata["documentCount"] == 1
    assert metadata["labelledCaseCount"] == 3
    assert metadata["excludedCaseCount"] == 0
    assert metadata["gitCommit"]
    assert metadata["pythonVersion"]
    assert metadata["platform"]
    # Timestamp is recorded but never part of pass/fail behavior.
    assert metadata["timestampUtc"]


def test_harness_fails_loudly_on_import_error(corpus: Path) -> None:
    (corpus.parent / "corpus" / "hist-001.docx").write_bytes(b"not a docx")
    with pytest.raises(RuntimeError, match="import failed"):
        run_validation(corpus)


def test_harness_fails_on_label_item_not_in_baseline(corpus: Path) -> None:
    labels_path = corpus.parent / "labels" / "hist-001.json"
    data = json.loads(labels_path.read_text(encoding="utf-8"))
    data["labels"].append(
        {
            "itemId": "item-ghost",
            "expectedResolution": "RESOLVED",
            "expectedStatus": "MISSING",
        }
    )
    labels_path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(RuntimeError, match="missing from"):
        run_validation(corpus)


def test_harness_output_is_json_serializable(corpus: Path) -> None:
    results = run_validation(corpus)
    text = json.dumps(results, ensure_ascii=False)
    assert "hist-001" in text


def test_harness_records_document_sanitization_state(corpus: Path) -> None:
    results = run_validation(corpus)
    assert results["documents"][0]["sanitizationState"] == "synthetic"


# ---------------------------------------------------------------------------
# Generic report rendering (T6.4 infrastructure)
# ---------------------------------------------------------------------------


def test_renderer_refuses_purely_synthetic_run(corpus: Path) -> None:
    from validation.metrics import compute_metrics
    from validation.report import render_report

    results = run_validation(corpus)
    with pytest.raises(ValueError, match="not historical evidence"):
        render_report(results, compute_metrics(results))


def test_renderer_rejects_missing_provenance() -> None:
    """A run with no documents (or no provenance) has no historical evidence."""
    from validation.metrics import compute_metrics
    from validation.report import is_historical_run, render_report

    assert not is_historical_run({"cases": [], "documents": []})
    assert not is_historical_run({"documents": [{"documentId": "x", "coverage": "COMPLETE"}]})
    with pytest.raises(ValueError, match="not historical evidence"):
        render_report(
            {"cases": [], "documents": []},
            compute_metrics({"cases": [], "documents": []}),
        )


def test_renderer_emits_measured_table_for_sanitized_run() -> None:
    from validation.metrics import compute_metrics
    from validation.report import render_report

    results = {
        "metadata": {
            "documentCount": 1,
            "ruleRevision": "task3-v1",
            "validationSchemaVersion": 1,
            "gitCommit": "abc123",
            "timestampUtc": "2026-09-23T00:00:00+00:00",
        },
        "baselineId": "baseline-1",
        "documents": [
            {
                "documentId": "hist-001",
                "coverage": "LIMITED",
                "coverageWarnings": ["tracked-revisions"],
                "sanitizationState": "sanitized",
            }
        ],
        "cases": [
            {
                "documentId": "hist-001",
                "itemId": "item-a",
                "includeInMetrics": True,
                "exclusionReason": "",
                "expected": {
                    "resolution": "RESOLVED",
                    "status": "CONFIGURED",
                    "comparison": None,
                    "strike": None,
                    "coverage": None,
                },
                "observed": {
                    "resolution": "RESOLVED",
                    "status": "MISSING",
                    "comparison": "NOT_COMPARED",
                    "strike": "NONE",
                    "coverage": "LIMITED",
                    "reviewReasons": [],
                    "comparisonReason": "",
                },
                "correct": {
                    "resolution": True,
                    "status": False,
                    "comparison": True,
                    "strike": True,
                },
                "discrepancyCategory": "",
                "notes": "",
            }
        ],
    }
    report = render_report(results, compute_metrics(results))
    assert "hist-001 / item-a" in report
    assert "RESOLVED/CONFIGURED" in report
    assert "RESOLVED/MISSING" in report
    assert "UNCLASSIFIED" in report
    assert "NOT MEASURED" in report
    assert "task3-v1" in report
    # Every metric row shows a denominator.
    assert "0/1 = 0.0%" in report
